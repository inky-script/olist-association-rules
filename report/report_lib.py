# -*- coding: utf-8 -*-
"""Thư viện định dạng báo cáo Word theo chuẩn trình bày đồ án (A4, Times New Roman 13,
giãn dòng 1.5, lề trái 3 cm). Dùng chung cho mọi phần của báo cáo.

Cú pháp nội dòng trong các hàm para()/bullet():
    **đậm**   *nghiêng*   `mã`   [[chỗ cần điền - tô vàng]]
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Times New Roman"
CODE_FONT = "Consolas"
HEADER_FILL = "D9E2F3"
CODE_FILL = "F2F2F2"
PLACEHOLDER_FILL = "FFF2CC"

INLINE = re.compile(r"(\*\*.+?\*\*|\[\[.+?\]\]|`.+?`|\*.+?\*)")


# ----------------------------------------------------------------------------
# Thiết lập tài liệu
# ----------------------------------------------------------------------------
def _set_font(style_or_run, name=FONT, size=None, bold=None, italic=None, color=None):
    font = style_or_run.font
    font.name = name
    if size:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if italic is not None:
        font.italic = italic
    if color is not None:
        font.color.rgb = color
    rpr = style_or_run.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)
    # Heading mặc định dùng theme font (asciiTheme...) sẽ ghi đè tên font -> xóa đi
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)


def new_document() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
    sec.left_margin, sec.right_margin = Cm(3.0), Cm(2.0)

    normal = doc.styles["Normal"]
    _set_font(normal, size=13)
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(6)
    pf.space_before = Pt(0)

    for name, size, italic, before in (("Heading 1", 15, False, 12),
                                       ("Heading 2", 13.5, False, 10),
                                       ("Heading 3", 13, True, 8)):
        st = doc.styles[name]
        _set_font(st, size=size, bold=True, italic=italic, color=RGBColor(0, 0, 0))
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.line_spacing = 1.3
        st.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number", "List Bullet 2"):
        _set_font(doc.styles[name], size=13)
    _set_font(doc.styles["Caption"], size=12, bold=False, italic=True, color=RGBColor(0, 0, 0))
    return doc


def add_page_number_footer(section, start_at=None, fmt="decimal"):
    """fmt: 'decimal' (1, 2, 3) hoặc 'lowerRoman' (i, ii, iii) cho phần mở đầu."""
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_field(p, "PAGE", "1")
    # section mới sao chép sectPr của section trước -> xóa pgNumType cũ rồi đặt lại
    for old in section._sectPr.findall(qn("w:pgNumType")):
        section._sectPr.remove(old)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:fmt"), fmt)
    if start_at is not None:
        pg.set(qn("w:start"), str(start_at))
    section._sectPr.append(pg)


def clear_footer(section):
    section.footer.is_linked_to_previous = False
    for p in section.footer.paragraphs:
        for r in p.runs:
            r.text = ""


def add_page_border(section):
    """Khung viền đôi cho trang bìa."""
    borders = OxmlElement("w:pgBorders")
    borders.set(qn("w:offsetFrom"), "page")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "thinThickSmallGap")
        el.set(qn("w:sz"), "24")
        el.set(qn("w:space"), "24")
        el.set(qn("w:color"), "1F3864")
        borders.append(el)
    # pgBorders phải đứng sau pgMar theo schema
    pgmar = section._sectPr.find(qn("w:pgMar"))
    pgmar.addnext(borders)


def new_section(doc):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    for b in sec._sectPr.findall(qn("w:pgBorders")):
        sec._sectPr.remove(b)
    return sec


def _add_field(paragraph, instr, placeholder=""):
    run = paragraph.add_run()
    for tag, attr in (("w:fldChar", "begin"),):
        el = OxmlElement(tag)
        el.set(qn("w:fldCharType"), attr)
        run._r.append(el)
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = f" {instr} "
    run._r.append(instr_el)
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run._r.append(sep)
    t = OxmlElement("w:t")
    t.text = placeholder
    run._r.append(t)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)
    return run


def add_toc(doc, title="MỤC LỤC"):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(p.add_run(title), size=15, bold=True)
    toc = doc.add_paragraph()
    _add_field(toc, 'TOC \\o "1-3" \\h \\z \\u',
               "Nhấp chuột phải → Update Field (hoặc nhấn F9) để cập nhật mục lục.")


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ----------------------------------------------------------------------------
# Văn bản
# ----------------------------------------------------------------------------
def add_inline(paragraph, text, size=None, base_bold=False, base_italic=False):
    for part in INLINE.split(text):
        if not part:
            continue
        bold, italic, code, placeholder = base_bold, base_italic, False, False
        if part.startswith("**") and part.endswith("**"):
            part, bold = part[2:-2], True
        elif part.startswith("[[") and part.endswith("]]"):
            part, placeholder = "[" + part[2:-2] + "]", True
        elif part.startswith("`") and part.endswith("`"):
            part, code = part[1:-1], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, italic = part[1:-1], True
        run = paragraph.add_run(part)
        _set_font(run, name=CODE_FONT if code else FONT,
                  size=(size - 1.5 if code and size else (11.5 if code else size)),
                  bold=bold, italic=italic)
        if placeholder:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return paragraph


def para(doc, text, indent=True, align="justify", size=None, bold=False, italic=False,
         space_after=None):
    p = doc.add_paragraph()
    p.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    if indent:
        p.paragraph_format.first_line_indent = Cm(1.0)
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    add_inline(p, text, size=size, base_bold=bold, base_italic=italic)
    return p


def bullet(doc, text, level=1, numbered=False):
    style = "List Number" if numbered else ("List Bullet" if level == 1 else "List Bullet 2")
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(3)
    add_inline(p, text)
    return p


def bullets(doc, items, numbered=False):
    for it in items:
        bullet(doc, it, numbered=numbered)


def h1(doc, text):
    return doc.add_heading(text.upper(), level=1)


def h2(doc, text):
    return doc.add_heading(text, level=2)


def h3(doc, text):
    return doc.add_heading(text, level=3)


def formula(doc, text):
    """Công thức đặt giữa dòng, chữ nghiêng."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    add_inline(p, text, base_italic=True)
    return p


# ----------------------------------------------------------------------------
# Bảng, hình, khối mã
# ----------------------------------------------------------------------------
def _shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def caption(doc, text):
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(10)
    add_inline(p, text, size=12, base_italic=True)
    return p


def table(doc, header, rows, widths_cm, cap=None, font_size=11.5, align=None):
    """align: danh sách 'l'/'c'/'r' cho từng cột."""
    if cap:
        # Chú thích bảng đặt PHÍA TRÊN bảng
        c = caption(doc, cap)
        c.paragraph_format.keep_with_next = True
        c.paragraph_format.space_after = Pt(3)
    align = align or ["l"] * len(header)
    amap = {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER,
            "r": WD_ALIGN_PARAGRAPH.RIGHT}
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, row_vals in enumerate([header] + rows):
        row = t.rows[i]
        if i == 0:
            _repeat_header(row)
        for j, val in enumerate(row_vals):
            cell = row.cells[j]
            cell.width = Cm(widths_cm[j])
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(2)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else amap[align[j]]
            add_inline(p, str(val), size=font_size, base_bold=(i == 0))
            if i == 0:
                _shade(cell, HEADER_FILL)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def code_block(doc, code, cap=None, font_size=9.5):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.rows[0].cells[0]
    cell.width = Cm(16.0)
    _shade(cell, CODE_FILL)
    lines = code.strip("\n").split("\n")
    first = cell.paragraphs[0]
    for k, line in enumerate(lines):
        p = first if k == 0 else cell.add_paragraph()
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(line if line else " ")
        _set_font(run, name=CODE_FONT, size=font_size)
    if cap:
        caption(doc, cap)
    else:
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def figure(doc, path, cap, width_cm=15.5, hint=None):
    """Chèn hình nếu tồn tại; nếu chưa có (chưa chạy code) thì chèn khung giữ chỗ tô vàng."""
    path = Path(path)
    if path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(path), width=Cm(width_cm))
    else:
        t = doc.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = t.rows[0].cells[0]
        cell.width = Cm(15.5)
        _shade(cell, PLACEHOLDER_FILL)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(24)
        p.paragraph_format.space_after = Pt(24)
        add_inline(p, f"[[CHÈN HÌNH: {hint or path.as_posix()}]]", size=12)
    caption(doc, cap)
