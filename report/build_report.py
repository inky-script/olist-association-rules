# -*- coding: utf-8 -*-
"""
Sinh báo cáo Word cho Đề tài 5 (Olist).
    python report/build_report.py            -> report/BaoCao_DoAn_Olist.docx

PHASE 4: trang bìa, phân công, mục 3–5. (PHASE 5 bổ sung mục 6–11 vào cùng file.)
Hình lấy từ outputs/figures/ nếu đã chạy Phase 2; nếu chưa có sẽ chèn khung giữ chỗ.
Các đoạn [trong ngoặc vuông, tô vàng] là số liệu cần điền từ output khi chạy dữ liệu thật.
"""
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from report_lib import (add_inline, add_page_border, add_page_number_footer, add_toc,
                        bullet, bullets, caption, clear_footer, code_block, figure, formula,
                        h1, h2, h3, new_document, new_section, page_break, para, table)

ROOT = Path(__file__).resolve().parents[1]
OUT_FIG = ROOT / "outputs" / "figures"
REPORT_FIG = ROOT / "report" / "figures"
OUT_FILE = ROOT / "report" / "BaoCao_DoAn_Olist.docx"

MEMBERS = [("Nguyễn Phan Đình Hào", "2402450074"),
           ("Nguyễn Thu Thảo", "2045240281"),
           ("Ngô Tuyết Vy", "2045240346")]


# ============================================================================
# Hình minh họa tự vẽ (sơ đồ quan hệ, pipeline)
# ============================================================================
def _box(ax, x, y, w, h, title, lines, color):
    ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.02,rounding_size=0.08",
                                fc="white", ec=color, lw=1.8))
    ax.add_patch(FancyBboxPatch((x, y + h - 0.42), w, 0.42,
                                boxstyle="round,pad=0.02,rounding_size=0.08", fc=color, ec=color))
    ax.text(x + w / 2, y + h - 0.21, title, ha="center", va="center", color="white",
            fontsize=11.5, fontweight="bold")
    for k, line in enumerate(lines):
        ax.text(x + 0.12, y + h - 0.7 - k * 0.3, line, ha="left", va="center", fontsize=10,
                family="DejaVu Sans Mono")


def _arrow(ax, p1, p2, label=None, color="#334155"):
    ax.add_patch(FancyArrowPatch(p1, p2, arrowstyle="-|>", mutation_scale=14, lw=1.4,
                                 color=color))
    if label:
        ax.text((p1[0] + p2[0]) / 2, (p1[1] + p2[1]) / 2 + 0.15, label, ha="center",
                fontsize=9.5, color=color, style="italic",
                bbox=dict(fc="white", ec="none", pad=1))


def draw_schema(path):
    fig, ax = plt.subplots(figsize=(10, 5.2))
    ax.set_xlim(0, 11)
    ax.set_ylim(0, 5.4)
    ax.axis("off")
    used, other = "#1F3864", "#94A3B8"
    _box(ax, 0.2, 2.6, 2.6, 2.5, "orders", ["PK order_id", "FK customer_id", "order_status",
                                            "order_purchase_ts", "..."], used)
    _box(ax, 4.1, 2.6, 2.7, 2.5, "order_items", ["PK order_id +", "   order_item_id",
                                                 "FK product_id", "FK seller_id", "price, freight"],
         used)
    _box(ax, 8.1, 2.6, 2.7, 2.5, "products", ["PK product_id", "FK product_category", "   _name",
                                              "weight, size...", "..."], used)
    _box(ax, 8.1, 0.15, 2.7, 1.6, "category_translation", ["PK product_category", "   _name",
                                                            "..._name_english"], used)
    _box(ax, 0.2, 0.15, 2.6, 1.6, "customers / payments", ["reviews (theo order_id)",
                                                            "geolocation (zip)"], other)
    _box(ax, 4.1, 0.15, 2.7, 1.6, "sellers", ["PK seller_id", "seller_city, state"], other)
    _arrow(ax, (2.8, 3.85), (4.1, 3.85), "1 : N  (order_id)")
    _arrow(ax, (6.8, 3.85), (8.1, 3.85), "N : 1  (product_id)")
    _arrow(ax, (9.45, 2.6), (9.45, 1.75), "N : 1")
    _arrow(ax, (5.45, 2.6), (5.45, 1.75), "N : 1", color=other)
    _arrow(ax, (1.5, 2.6), (1.5, 1.75), "1 : N", color=other)
    ax.text(5.5, 5.3, "Màu đậm: 4 bảng dùng trong đề tài  ·  Màu xám: bảng không sử dụng",
            ha="center", fontsize=10.5, color="#334155")
    fig.tight_layout()
    fig.savefig(path, dpi=200)
    plt.close(fig)


def draw_pipeline(path):
    steps = [("1. Nạp dữ liệu", "4 file CSV\nutf-8-sig"),
             ("2. Kiểm tra khóa", "PK duy nhất\nFK mồ côi"),
             ("3. Ghép bảng", "many_to_one\nassert số dòng"),
             ("4. Xử lý thiếu", "unknown,\ndịch bổ sung"),
             ("5. Lọc & ngoại lệ", "bỏ canceled /\nunavailable; IQR"),
             ("6. Transaction", "groupby\norder_id → set"),
             ("7. One-hot", "TransactionEncoder\nma trận bool")]
    fig, ax = plt.subplots(figsize=(11, 2.8))
    ax.set_xlim(0, 12.2)
    ax.set_ylim(0, 2.6)
    ax.axis("off")
    w, gap = 1.5, 0.22
    for k, (title, sub) in enumerate(steps):
        x = 0.1 + k * (w + gap)
        color = "#1F3864" if k < 5 else "#C55A11"
        ax.add_patch(FancyBboxPatch((x, 0.45), w, 1.7, boxstyle="round,pad=0.02,rounding_size=0.1",
                                    fc=color, ec=color))
        ax.text(x + w / 2, 1.75, title, ha="center", va="center", color="white", fontsize=10,
                fontweight="bold")
        ax.text(x + w / 2, 1.05, sub, ha="center", va="center", color="white", fontsize=9)
        if k < len(steps) - 1:
            _arrow(ax, (x + w, 1.3), (x + w + gap, 1.3))
    ax.text(0.1 + 2.5 * (w + gap), 0.15, "Làm sạch & ghép bảng (Phase 1)", ha="center",
            fontsize=9, color="#1F3864")
    ax.text(0.1 + 5.5 * (w + gap) + 0.4, 0.15, "Biểu diễn giao dịch (Phase 2)", ha="center",
            fontsize=9, color="#C55A11")
    fig.tight_layout()
    fig.savefig(path, dpi=200)
    plt.close(fig)


# ============================================================================
# 1. TRANG BÌA
# ============================================================================
def cover(doc):
    sec = doc.sections[0]
    add_page_border(sec)
    clear_footer(sec)

    def line(text, size=13, bold=False, before=0, after=0, italic=False):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.line_spacing = 1.2
        add_inline(p, text, size=size, base_bold=bold, base_italic=italic)

    line("BỘ CÔNG THƯƠNG", 14, True)
    line("TRƯỜNG ĐẠI HỌC CÔNG THƯƠNG THÀNH PHỐ HỒ CHÍ MINH", 14, True)
    line("KHOA CÔNG NGHỆ THÔNG TIN", 14, True)
    line("———————o0o———————", 12, after=18)
    line("[[Chèn logo HUIT tại đây]]", 12, before=24, after=36)
    line("ĐỒ ÁN MÔN HỌC", 20, True, after=4)
    line("THỰC HÀNH KHAI THÁC DỮ LIỆU", 18, True, after=30)
    line("ĐỀ TÀI SỐ 5", 15, True, after=6)
    line("KHAI THÁC QUAN HỆ MUA KÈM GIỮA CÁC DANH MỤC SẢN PHẨM", 16, True)
    line("TRONG THƯƠNG MẠI ĐIỆN TỬ OLIST", 16, True, after=48)

    t = doc.add_table(rows=0, cols=2)
    rows = [("Giảng viên hướng dẫn:", "[[ThS./TS. họ tên giảng viên]]"),
            ("Lớp học phần:", "[[mã lớp]]"),
            ("Nhóm sinh viên thực hiện:", "")]
    rows += [(f"   {i}. {name}", f"MSSV: {mssv}") for i, (name, mssv) in enumerate(MEMBERS, 1)]
    for left, right in rows:
        cells = t.add_row().cells
        for cell, text, bold in ((cells[0], left, left.endswith(":")), (cells[1], right, False)):
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.3
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, text, size=13.5, base_bold=bold)
    from docx.shared import Cm
    for row in t.rows:
        row.cells[0].width, row.cells[1].width = Cm(7.0), Cm(7.5)
    t.alignment = 1

    line("TP. Hồ Chí Minh, tháng [[..]] năm 2026", 13, italic=True, before=70)


# ============================================================================
# 2. BẢNG PHÂN CÔNG
# ============================================================================
def assignment(doc):
    h1(doc, "2. Bảng phân công công việc và tỷ lệ đóng góp")
    para(doc, "Nhóm gồm 3 thành viên, cùng thống nhất phương pháp ngay từ đầu, chia việc theo "
              "luồng xử lý dữ liệu và rà soát chéo sản phẩm của nhau trước khi tổng hợp. "
              "Bảng 2.1 trình bày nhiệm vụ cụ thể và tỷ lệ đóng góp của từng thành viên.")
    table(doc, ["STT", "Thành viên", "Nhiệm vụ chính", "Mức độ hoàn thành", "Tỷ lệ đóng góp"],
          [["1", "Thành viên 1",
            "Tìm hiểu và mô tả dữ liệu Olist; kiểm tra khóa, ghép bảng, xử lý giá trị thiếu và "
            "ngoại lệ; xây dựng transaction và one-hot encoding; viết mục 3, 4, 5.",
            "Hoàn thành tốt", "100%"],
           ["2", "Thành viên 2",
            "Cài đặt Apriori và FP-Growth bằng mlxtend; thiết kế thực nghiệm khảo sát minsup "
            "(runtime, số tập phổ biến, số luật); thí nghiệm mở rộng mức sản phẩm; vẽ biểu đồ "
            "so sánh; viết mục 6, 7.", "Hoàn thành tốt", "100%"],
           ["3", "Thành viên 3",
            "Lọc và diễn giải các luật tiêu biểu, đề xuất chiến lược cross-selling; xây dựng "
            "ứng dụng demo Streamlit; viết mục 8, 9, 10, 11; tổng hợp, định dạng báo cáo.",
            "Hoàn thành tốt", "100%"]],
          [1.3, 2.5, 7.2, 2.3, 2.3], cap="Bảng 2.1: Phân công công việc và tỷ lệ đóng góp",
          align=["c", "l", "l", "c", "c"])
    para(doc, "Ngoài phần việc chính, cả nhóm cùng tham gia họp định kỳ, thảo luận lựa chọn "
              "tham số, kiểm thử mã nguồn và đọc soát toàn văn báo cáo. Các thành viên cam kết "
              "nội dung trong báo cáo là kết quả làm việc của nhóm; mọi tài liệu và mã nguồn "
              "tham khảo đều được trích dẫn tại mục Tài liệu tham khảo.")


def abbreviations(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(p, "DANH MỤC TỪ VIẾT TẮT", size=15, base_bold=True)
    table(doc, ["Từ viết tắt", "Tiếng Anh", "Ý nghĩa"],
          [["TMĐT", "E-commerce", "Thương mại điện tử"],
           ["MBA", "Market Basket Analysis", "Phân tích giỏ hàng"],
           ["ARM", "Association Rule Mining", "Khai thác luật kết hợp"],
           ["FP-Growth", "Frequent Pattern Growth", "Thuật toán tăng trưởng mẫu phổ biến"],
           ["FP-tree", "Frequent Pattern tree", "Cây mẫu phổ biến"],
           ["minsup", "Minimum support", "Ngưỡng độ hỗ trợ tối thiểu"],
           ["minconf", "Minimum confidence", "Ngưỡng độ tin cậy tối thiểu"],
           ["EDA", "Exploratory Data Analysis", "Phân tích khám phá dữ liệu"],
           ["PK / FK", "Primary key / Foreign key", "Khóa chính / Khóa ngoại"],
           ["IQR", "Interquartile range", "Khoảng tứ phân vị"],
           ["BRL", "Brazilian Real", "Đồng Real Brazil"]],
          [3.0, 5.5, 7.0])


# ============================================================================
# 3. GIỚI THIỆU BÀI TOÁN
# ============================================================================
def section3(doc):
    h1(doc, "3. Giới thiệu bài toán")

    h2(doc, "3.1. Bối cảnh")
    para(doc, "Thương mại điện tử (TMĐT) đã trở thành một kênh bán lẻ chủ lực trên toàn cầu. "
              "Tại Brazil, quốc gia đông dân nhất Mỹ Latinh, mua sắm trực tuyến tăng trưởng mạnh "
              "trong thập kỷ qua, kéo theo sự xuất hiện của nhiều sàn giao dịch và nền tảng trung "
              "gian. **Olist** là một công ty khởi nghiệp có trụ sở tại Curitiba, hoạt động như "
              "một “cửa hàng bách hóa” trên các sàn TMĐT lớn: các nhà bán lẻ vừa và nhỏ đăng ký "
              "với Olist, sau đó sản phẩm của họ được niêm yết và bán trên các sàn đối tác thông "
              "qua một hợp đồng duy nhất, còn Olist đảm nhận phần kết nối, vận hành và hậu cần.")
    para(doc, "Năm 2018, Olist công bố trên Kaggle bộ dữ liệu *Brazilian E-Commerce Public "
              "Dataset by Olist*, gồm khoảng 100 nghìn đơn hàng thực tế phát sinh từ tháng "
              "9/2016 đến tháng 10/2018, đã được ẩn danh hóa. Đây là một trong số ít bộ dữ liệu "
              "TMĐT công khai có cấu trúc quan hệ đầy đủ (đơn hàng, sản phẩm, người bán, thanh "
              "toán, đánh giá, vị trí địa lý), vì vậy rất phù hợp để thực hành các kỹ thuật khai "
              "thác dữ liệu trên một bài toán kinh doanh sát thực tế.")
    para(doc, "Một trong những câu hỏi then chốt của mọi nhà bán lẻ là: *khách hàng thường mua "
              "những gì cùng nhau?* Trả lời được câu hỏi này giúp doanh nghiệp thiết kế các hoạt "
              "động **bán kèm (cross-selling)** như gợi ý “thường được mua cùng”, đóng gói combo "
              "sản phẩm, đặt ưu đãi cho món thứ hai hay sắp xếp vị trí sản phẩm trên trang. Kỹ "
              "thuật kinh điển để giải quyết vấn đề này là **phân tích giỏ hàng (Market Basket "
              "Analysis – MBA)** dựa trên **khai thác luật kết hợp (Association Rule Mining)**, "
              "do Agrawal và cộng sự đề xuất từ năm 1993.")
    para(doc, "Khác với một siêu thị truyền thống, dữ liệu Olist có hai đặc thù quan trọng. Thứ "
              "nhất, số lượng sản phẩm rất lớn (khoảng 33 nghìn mã) nhưng phần lớn chỉ được bán "
              "vài lần, nên khai thác trực tiếp ở mức sản phẩm sẽ cho ra rất ít mẫu có ý nghĩa "
              "thống kê. Thứ hai, đa số đơn hàng chỉ chứa một sản phẩm, nên giỏ hàng cực kỳ thưa. "
              "Vì vậy đề tài lựa chọn khai thác ở **mức danh mục sản phẩm** (khoảng 70 danh mục), "
              "là mức trừu tượng đủ khái quát để thu được các quan hệ mua kèm ổn định và vẫn đủ "
              "cụ thể để đề xuất hành động kinh doanh.")

    h2(doc, "3.2. Phát biểu bài toán")
    para(doc, "Gọi *I = {i₁, i₂, …, iₘ}* là tập các danh mục sản phẩm (item) và "
              "*D = {t₁, t₂, …, tₙ}* là cơ sở dữ liệu giao dịch, trong đó mỗi giao dịch "
              "*tⱼ ⊆ I* là tập các danh mục **khác nhau** xuất hiện trong cùng một đơn hàng "
              "(xác định bởi `order_id`). Một **luật kết hợp** có dạng *X → Y*, với *X, Y ⊆ I*, "
              "*X ≠ ∅, Y ≠ ∅* và *X ∩ Y = ∅*; trong đó *X* gọi là tiền đề (antecedent), *Y* là "
              "hệ quả (consequent).")
    para(doc, "Bài toán đặt ra là tìm tất cả các luật *X → Y* thỏa đồng thời ba điều kiện:")
    formula(doc, "support(X ∪ Y) ≥ minsup,     confidence(X → Y) ≥ minconf,     lift(X → Y) ≥ minlift")
    para(doc, "Từ tập luật thu được, chọn ra khoảng 10–20 luật có giá trị nhất (lift rõ rệt lớn "
              "hơn 1 và có đủ số đơn hàng hỗ trợ), diễn giải ý nghĩa kinh doanh và xây dựng một "
              "ứng dụng gợi ý mua kèm. Bài toán được tóm tắt trong Bảng 3.1.")
    table(doc, ["Thành phần", "Mô tả"],
          [["Đầu vào", "Bảng `order_items`, `products`, `product_category_name_translation` và "
                       "`orders` của bộ dữ liệu Olist."],
           ["Đơn vị giao dịch", "Một đơn hàng (`order_id`)."],
           ["Item", "Danh mục sản phẩm, dùng tên tiếng Anh (`product_category_name_english`)."],
           ["Thuật toán", "Apriori và FP-Growth (thư viện `mlxtend`)."],
           ["Đầu ra", "Tập phổ biến, tập luật kèm Support, Confidence, Lift; Top 10–20 luật; "
                      "bảng so sánh hiệu năng; ứng dụng Streamlit gợi ý mua kèm."]],
          [3.8, 11.7], cap="Bảng 3.1: Tóm tắt đầu vào – đầu ra của bài toán")

    h2(doc, "3.3. Mục tiêu đề tài")
    bullets(doc, [
        "**Mục tiêu dữ liệu:** xây dựng quy trình tiền xử lý có kiểm soát, ghép đúng khóa giữa "
        "`order_items` – `products` – bảng dịch – `orders`, xử lý giá trị thiếu và ngoại lệ, "
        "chuyển dữ liệu về dạng giao dịch và ma trận one-hot.",
        "**Mục tiêu thuật toán:** cài đặt Apriori và FP-Growth, kiểm chứng hai thuật toán cho "
        "cùng tập phổ biến, khảo sát sự thay đổi của thời gian chạy, số tập phổ biến và số luật "
        "khi thay đổi ngưỡng minsup.",
        "**Mục tiêu tri thức:** đánh giá luật bằng Support, Confidence, Lift; loại các luật có "
        "lift xấp xỉ 1; chọn 10–20 luật tiêu biểu và diễn giải ý nghĩa cross-selling.",
        "**Mục tiêu ứng dụng:** xây dựng ứng dụng Streamlit cho phép chọn một danh mục (hoặc một "
        "giỏ hàng) và nhận gợi ý các danh mục nên bán kèm.",
    ])

    h2(doc, "3.4. Câu hỏi khai thác")
    para(doc, "Để định hướng quá trình phân tích, nhóm đặt ra năm câu hỏi khai thác cụ thể:")
    bullets(doc, [
        "**Q1.** Mức độ “mua kèm” trong dữ liệu Olist thực tế cao đến đâu, tức là bao nhiêu phần "
        "trăm đơn hàng chứa từ hai danh mục trở lên?",
        "**Q2.** Những cặp (hoặc bộ ba) danh mục nào thường xuất hiện cùng nhau với mức liên kết "
        "vượt trội so với ngẫu nhiên (lift > 1)?",
        "**Q3.** Với mỗi danh mục, nên gợi ý bán kèm những danh mục nào và mức độ tin cậy ra sao?",
        "**Q4.** Khi giảm ngưỡng minsup, số tập phổ biến, số luật và thời gian chạy thay đổi thế "
        "nào; Apriori hay FP-Growth hiệu quả hơn trong bối cảnh dữ liệu này?",
        "**Q5.** Kết quả khai thác có thể chuyển thành những hành động cross-selling cụ thể nào "
        "cho nhà bán hàng trên Olist?",
    ], numbered=False)

    h2(doc, "3.5. Ý nghĩa và giá trị")
    para(doc, "**Về mặt kinh doanh**, các luật mua kèm là cơ sở định lượng cho việc thiết kế "
              "combo sản phẩm, hệ thống gợi ý “khách hàng thường mua cùng”, chương trình khuyến "
              "mãi chéo danh mục và bố cục trang danh mục. Những hoạt động này giúp tăng giá trị "
              "trung bình mỗi đơn hàng, đồng thời có thể giảm chi phí vận chuyển trên mỗi sản "
              "phẩm khi khách mua nhiều món trong cùng một đơn.")
    para(doc, "**Về mặt học thuật**, đề tài cho phép đối chiếu lý thuyết với thực nghiệm: so sánh "
              "hai thuật toán khai thác tập phổ biến kinh điển trên dữ liệu thật, quan sát tác "
              "động của độ thưa dữ liệu và của kích thước không gian item lên hiệu năng, qua đó "
              "hiểu rõ khi nào nên dùng thuật toán nào.")
    para(doc, "**Về mặt kỹ năng**, nhóm rèn luyện được toàn bộ quy trình của một dự án khai thác "
              "dữ liệu: hiểu nghiệp vụ, hiểu dữ liệu, tiền xử lý, mô hình hóa, đánh giá và triển "
              "khai thành sản phẩm demo.")

    h2(doc, "3.6. Phạm vi và giới hạn")
    bullets(doc, [
        "Chỉ sử dụng 4 bảng: `orders`, `order_items`, `products`, "
        "`product_category_name_translation`; các bảng khách hàng, thanh toán, đánh giá, người "
        "bán, vị trí địa lý nằm ngoài phạm vi.",
        "Khai thác chính ở mức danh mục; mức sản phẩm (`product_id`) chỉ dùng cho thí nghiệm "
        "so sánh hiệu năng thuật toán.",
        "Luật được khai thác trong phạm vi **một đơn hàng**; chưa xét hành vi mua lặp lại của "
        "cùng khách hàng theo thời gian (luật tuần tự).",
        "Dữ liệu thuộc giai đoạn 2016–2018 tại Brazil, nên các kết luận mang tính tham khảo khi "
        "áp dụng cho thị trường khác.",
    ])

    h2(doc, "3.7. Quy trình thực hiện")
    para(doc, "Đề tài được triển khai theo tinh thần của quy trình CRISP-DM, gồm sáu bước nối "
              "tiếp: (1) hiểu bài toán kinh doanh; (2) hiểu dữ liệu và đánh giá chất lượng; "
              "(3) tiền xử lý, ghép bảng, tạo giao dịch; (4) khai thác tập phổ biến và luật kết "
              "hợp bằng Apriori và FP-Growth; (5) đánh giá, so sánh và diễn giải luật; (6) triển "
              "khai ứng dụng demo bằng Streamlit. Mã nguồn được tổ chức theo đúng các bước này "
              "(Bảng 3.2) để thuận tiện tái lập kết quả.")
    table(doc, ["Tệp mã nguồn", "Chức năng", "Đầu ra chính"],
          [["`src/phase1_preprocessing.py`", "Nạp, kiểm tra khóa, ghép bảng, xử lý thiếu, "
            "lọc trạng thái, gắn cờ ngoại lệ", "`order_items_clean.csv`, báo cáo chất lượng"],
           ["`src/phase2_mining.py`", "Tạo giao dịch, one-hot, Apriori và FP-Growth, "
            "benchmark, sinh và lọc luật", "Bảng benchmark, Top 20 luật, biểu đồ, `rules.csv`"],
           ["`app.py`", "Ứng dụng Streamlit gợi ý mua kèm", "Giao diện web tương tác"]],
          [5.0, 5.8, 4.7], cap="Bảng 3.2: Tổ chức mã nguồn theo quy trình")


# ============================================================================
# 4. DỮ LIỆU
# ============================================================================
def section4(doc):
    h1(doc, "4. Dữ liệu")

    h2(doc, "4.1. Nguồn dữ liệu")
    para(doc, "Bộ dữ liệu *Brazilian E-Commerce Public Dataset by Olist* được Olist công bố trên "
              "nền tảng Kaggle tại địa chỉ https://www.kaggle.com/datasets/olistbr/brazilian-ecommerce, "
              "theo giấy phép CC BY-NC-SA 4.0 (được phép sử dụng cho mục đích phi thương mại như "
              "học tập, nghiên cứu, với điều kiện ghi nguồn). Dữ liệu là các đơn hàng thực tế đã "
              "được ẩn danh: mã khách hàng, mã người bán, mã sản phẩm đều được thay bằng chuỗi "
              "băm; tên công ty đối tác trong phần đánh giá được thay bằng tên các nhà Đại Hùng "
              "Tinh trong *Game of Thrones*.")
    para(doc, "Bộ dữ liệu gồm 9 tệp CSV liên kết với nhau theo mô hình quan hệ, tập trung "
              "quanh bảng `orders`. Các đơn hàng trải dài từ tháng 9/2016 đến tháng 10/2018, "
              "phát sinh trên nhiều sàn TMĐT tại Brazil.")

    h2(doc, "4.2. Tổng quan các bảng dữ liệu")
    table(doc, ["Bảng (tệp CSV)", "Số dòng", "Số cột", "Nội dung", "Sử dụng"],
          [["olist_orders_dataset", "99.441", "8", "Thông tin đơn hàng, trạng thái, mốc thời gian", "✔"],
           ["olist_order_items_dataset", "112.650", "7", "Từng sản phẩm trong đơn, giá, phí vận chuyển", "✔"],
           ["olist_products_dataset", "32.951", "9", "Thuộc tính sản phẩm, danh mục (tiếng Bồ Đào Nha)", "✔"],
           ["product_category_name_translation", "71", "2", "Bảng dịch tên danh mục sang tiếng Anh", "✔"],
           ["olist_customers_dataset", "99.441", "5", "Khách hàng, mã bưu chính, thành phố", "–"],
           ["olist_order_payments_dataset", "103.886", "5", "Phương thức, số kỳ, giá trị thanh toán", "–"],
           ["olist_order_reviews_dataset", "99.224", "7", "Điểm và nội dung đánh giá", "–"],
           ["olist_sellers_dataset", "3.095", "4", "Người bán và vị trí", "–"],
           ["olist_geolocation_dataset", "1.000.163", "5", "Tọa độ theo mã bưu chính", "–"]],
          [5.4, 2.0, 1.4, 5.5, 1.3], cap="Bảng 4.1: Danh sách 9 bảng của bộ dữ liệu Olist",
          align=["l", "r", "c", "l", "c"], font_size=11)
    para(doc, "Trong phạm vi đề tài, nhóm sử dụng bốn bảng được đánh dấu ✔. Bảng `order_items` "
              "là bảng trung tâm vì mỗi dòng tương ứng với một sản phẩm trong một đơn hàng; bảng "
              "`products` cung cấp danh mục của sản phẩm; bảng dịch giúp chuẩn hóa tên danh mục "
              "sang tiếng Anh để dễ diễn giải; bảng `orders` cung cấp trạng thái đơn hàng phục vụ "
              "việc lọc các đơn không hợp lệ.")

    h2(doc, "4.3. Lược đồ quan hệ giữa các bảng")
    schema = REPORT_FIG / "schema.png"
    figure(doc, schema, "Hình 4.1: Lược đồ quan hệ giữa các bảng dữ liệu Olist", width_cm=15.5)
    para(doc, "Như Hình 4.1, quan hệ giữa `orders` và `order_items` là một – nhiều qua khóa "
              "`order_id`: một đơn hàng có thể gồm nhiều dòng sản phẩm, đánh số bởi "
              "`order_item_id` (1, 2, 3, …). Quan hệ giữa `order_items` và `products` là nhiều – "
              "một qua `product_id`; giữa `products` và bảng dịch là nhiều – một qua "
              "`product_category_name`. Khóa chính của `order_items` là cặp "
              "(`order_id`, `order_item_id`). Điểm cần lưu ý: nếu khách mua 3 đơn vị cùng một sản "
              "phẩm, bảng `order_items` sẽ có 3 dòng với cùng `product_id`, chứ không có cột số lượng.")

    h2(doc, "4.4. Mô tả chi tiết các bảng sử dụng")
    h3(doc, "4.4.1. Bảng order_items")
    table(doc, ["Thuộc tính", "Kiểu", "Mô tả"],
          [["order_id", "string", "Mã đơn hàng (khóa ngoại tới `orders`)"],
           ["order_item_id", "int", "Số thứ tự sản phẩm trong đơn (1, 2, …)"],
           ["product_id", "string", "Mã sản phẩm (khóa ngoại tới `products`)"],
           ["seller_id", "string", "Mã người bán"],
           ["shipping_limit_date", "datetime", "Hạn chót người bán giao hàng cho đơn vị vận chuyển"],
           ["price", "float", "Giá sản phẩm (BRL)"],
           ["freight_value", "float", "Phí vận chuyển phân bổ cho sản phẩm (BRL)"]],
          [4.2, 2.3, 9.0], cap="Bảng 4.2: Mô tả thuộc tính bảng order_items")
    h3(doc, "4.4.2. Bảng products")
    table(doc, ["Thuộc tính", "Kiểu", "Mô tả"],
          [["product_id", "string", "Mã sản phẩm (khóa chính)"],
           ["product_category_name", "string", "Tên danh mục bằng tiếng Bồ Đào Nha"],
           ["product_name_lenght", "float", "Số ký tự của tên sản phẩm (tên cột sai chính tả gốc)"],
           ["product_description_lenght", "float", "Số ký tự phần mô tả"],
           ["product_photos_qty", "float", "Số ảnh sản phẩm"],
           ["product_weight_g", "float", "Khối lượng (gram)"],
           ["product_length_cm / height_cm / width_cm", "float", "Kích thước (cm)"]],
          [5.6, 2.0, 7.9], cap="Bảng 4.3: Mô tả thuộc tính bảng products")
    h3(doc, "4.4.3. Bảng product_category_name_translation")
    table(doc, ["Thuộc tính", "Kiểu", "Mô tả"],
          [["product_category_name", "string", "Tên danh mục tiếng Bồ Đào Nha (khóa chính)"],
           ["product_category_name_english", "string", "Tên danh mục tiếng Anh tương ứng"]],
          [6.6, 1.8, 7.1], cap="Bảng 4.4: Mô tả thuộc tính bảng dịch danh mục")
    h3(doc, "4.4.4. Bảng orders")
    table(doc, ["Thuộc tính", "Kiểu", "Mô tả"],
          [["order_id", "string", "Mã đơn hàng (khóa chính)"],
           ["customer_id", "string", "Mã khách hàng theo đơn"],
           ["order_status", "string", "Trạng thái: delivered, shipped, canceled, …"],
           ["order_purchase_timestamp", "datetime", "Thời điểm đặt hàng"],
           ["order_approved_at", "datetime", "Thời điểm duyệt thanh toán"],
           ["order_delivered_carrier_date", "datetime", "Thời điểm giao cho đơn vị vận chuyển"],
           ["order_delivered_customer_date", "datetime", "Thời điểm khách nhận hàng"],
           ["order_estimated_delivery_date", "datetime", "Ngày giao dự kiến"]],
          [5.6, 2.0, 7.9], cap="Bảng 4.5: Mô tả thuộc tính bảng orders")

    h2(doc, "4.5. Đánh giá chất lượng dữ liệu")
    para(doc, "Trước khi tiền xử lý, nhóm thực hiện kiểm tra chất lượng dữ liệu một cách hệ thống "
              "bằng hàm `quality_report()` và `check_keys()` trong Phase 1. Kết quả được tổng "
              "hợp theo năm khía cạnh: giá trị thiếu, tính nhất quán của khóa, tính đầy đủ của "
              "bảng dịch, trạng thái đơn hàng và giá trị ngoại lệ.")
    h3(doc, "4.5.1. Giá trị thiếu")
    table(doc, ["Bảng", "Thuộc tính", "Số giá trị thiếu", "Ảnh hưởng đến đề tài"],
          [["products", "product_category_name", "610", "**Có** – không xác định được danh mục"],
           ["products", "product_name_lenght, product_description_lenght, product_photos_qty",
            "610", "Không – không dùng các cột này"],
           ["products", "product_weight_g, kích thước", "2", "Không"],
           ["orders", "order_approved_at", "160", "Không"],
           ["orders", "order_delivered_carrier_date", "1.783", "Không"],
           ["orders", "order_delivered_customer_date", "2.965", "Không"],
           ["order_items", "(tất cả)", "0", "–"],
           ["translation", "(tất cả)", "0", "–"]],
          [2.4, 6.0, 2.6, 4.5], cap="Bảng 4.6: Thống kê giá trị thiếu",
          align=["l", "l", "r", "l"], font_size=11)
    para(doc, "Có 610 sản phẩm (khoảng 1,85% số sản phẩm) không có danh mục. Đây là loại thiếu "
              "duy nhất ảnh hưởng trực tiếp đến bài toán, vì danh mục chính là item cần khai "
              "thác. Các giá trị thiếu ở bảng `orders` chủ yếu rơi vào những đơn chưa giao hoặc "
              "đã hủy, nên không ảnh hưởng tới việc xây dựng giao dịch.")
    h3(doc, "4.5.2. Tính nhất quán của khóa")
    para(doc, "Kết quả kiểm tra cho thấy: `orders.order_id`, `products.product_id` và "
              "`translation.product_category_name` đều là khóa duy nhất; cặp (`order_id`, "
              "`order_item_id`) là duy nhất trong `order_items`; mọi `product_id` và `order_id` "
              "trong `order_items` đều tồn tại trong bảng cha tương ứng (không có khóa ngoại “mồ "
              "côi”). Điều này đảm bảo phép ghép bảng không làm nhân bản hay mất dòng. Ngoài ra, "
              "khoảng 775 đơn hàng trong `orders` không có dòng nào trong `order_items` "
              "(chủ yếu là đơn bị hủy hoặc không khả dụng); các đơn này tự động bị loại khi lấy "
              "`order_items` làm bảng gốc.")
    h3(doc, "4.5.3. Tính đầy đủ của bảng dịch")
    para(doc, "Bảng `products` có 73 danh mục khác nhau trong khi bảng dịch chỉ có 71 dòng. Hai "
              "danh mục không có bản dịch là `pc_gamer` và "
              "`portateis_cozinha_e_preparadores_de_alimentos` (đồ dùng nhà bếp cầm tay và máy "
              "chế biến thực phẩm). Ngoài ra, tệp dịch được lưu với ký tự BOM ở đầu, nếu đọc "
              "bằng mã hóa mặc định thì tên cột đầu tiên sẽ bị sai thành "
              "`\\ufeffproduct_category_name` và phép ghép sẽ thất bại.")
    h3(doc, "4.5.4. Trạng thái đơn hàng")
    table(doc, ["Trạng thái", "Ý nghĩa", "Số đơn", "Tỷ lệ"],
          [["delivered", "Đã giao thành công", "96.478", "97,02%"],
           ["shipped", "Đang vận chuyển", "1.107", "1,11%"],
           ["canceled", "Đã hủy", "625", "0,63%"],
           ["unavailable", "Không khả dụng", "609", "0,61%"],
           ["invoiced", "Đã xuất hóa đơn", "314", "0,32%"],
           ["processing", "Đang xử lý", "301", "0,30%"],
           ["created", "Vừa tạo", "5", "0,01%"],
           ["approved", "Đã duyệt", "2", "0,00%"],
           ["**Tổng**", "", "**99.441**", "**100%**"]],
          [3.0, 5.5, 3.0, 2.5], cap="Bảng 4.7: Phân bố trạng thái đơn hàng trong bảng orders",
          align=["l", "l", "r", "r"])
    para(doc, "Hơn 97% đơn hàng đã giao thành công. Các trạng thái `canceled` và `unavailable` "
              "phản ánh giao dịch không hoàn tất, nên không thể hiện đúng hành vi mua kèm của "
              "khách hàng.")
    h3(doc, "4.5.5. Giá trị ngoại lệ")
    para(doc, "Giá sản phẩm (`price`) dao động từ 0,85 đến 6.735 BRL, giá trung bình khoảng "
              "120 BRL; phân phối lệch phải mạnh với một đuôi dài các sản phẩm giá trị cao. Phí "
              "vận chuyển (`freight_value`) dao động từ 0 đến khoảng 410 BRL. Không có giá âm "
              "hay giá bằng 0. Các giá trị lớn là hợp lệ về nghiệp vụ (ví dụ đồ điện tử, nội "
              "thất cao cấp) chứ không phải lỗi nhập liệu.")

    h2(doc, "4.6. Nhận xét chung")
    para(doc, "Dữ liệu Olist có chất lượng tương đối tốt: khóa nhất quán, bảng `order_items` "
              "không có giá trị thiếu. Các vấn đề cần xử lý gồm: (1) sản phẩm thiếu danh mục, "
              "(2) hai danh mục chưa có bản dịch và lỗi BOM của tệp dịch, (3) các đơn hàng có "
              "trạng thái không hợp lệ, (4) phân phối giá lệch. Ngoài ra, đặc điểm quan trọng "
              "nhất đối với bài toán là **tính thưa của giỏ hàng**: phần lớn đơn hàng chỉ chứa "
              "một sản phẩm. Đặc điểm này chi phối trực tiếp việc chọn ngưỡng minsup và được "
              "phân tích chi tiết ở mục 5.9.")


# ============================================================================
# 5. TIỀN XỬ LÝ DỮ LIỆU
# ============================================================================
def section5(doc):
    h1(doc, "5. Tiền xử lý dữ liệu")

    h2(doc, "5.1. Quy trình tổng quát")
    pipeline = REPORT_FIG / "pipeline.png"
    figure(doc, pipeline, "Hình 5.1: Quy trình tiền xử lý dữ liệu", width_cm=16)
    para(doc, "Quy trình tiền xử lý gồm bảy bước (Hình 5.1). Năm bước đầu thực hiện trong "
              "`phase1_preprocessing.py`, nhằm tạo ra bảng chi tiết đơn hàng đã làm sạch "
              "`order_items_clean.csv`. Hai bước cuối thực hiện trong `phase2_mining.py`, nhằm "
              "chuyển bảng này thành biểu diễn giao dịch phù hợp với thuật toán khai thác luật "
              "kết hợp. Nguyên tắc xuyên suốt là **mọi thao tác đều được kiểm chứng bằng mã** "
              "(assert, đối chiếu số dòng), chứ không chỉ quan sát thủ công.")

    h2(doc, "5.2. Nạp dữ liệu")
    para(doc, "Các tệp CSV được đọc bằng `pandas`, trong đó các cột thời gian được chuyển sang "
              "kiểu `datetime` ngay khi đọc. Riêng tệp dịch được đọc với tham số "
              "`encoding=\"utf-8-sig\"` để loại bỏ ký tự BOM ở đầu tệp. Chương trình kiểm tra sự "
              "tồn tại của cả 4 tệp trước khi chạy và báo lỗi rõ ràng nếu thiếu.")
    code_block(doc, '''
dfs = {
    "orders": pd.read_csv(data_dir / FILES["orders"], parse_dates=date_cols),
    "order_items": pd.read_csv(data_dir / FILES["order_items"],
                               parse_dates=["shipping_limit_date"]),
    "products": pd.read_csv(data_dir / FILES["products"]),
    # utf-8-sig: file dịch có BOM ở header
    "translation": pd.read_csv(data_dir / FILES["translation"], encoding="utf-8-sig"),
}''', cap="Đoạn mã 5.1: Nạp dữ liệu và xử lý BOM")

    h2(doc, "5.3. Ghép bảng theo khóa")
    para(doc, "Bảng `order_items` được chọn làm bảng gốc vì đây là nơi ghi nhận từng sản phẩm "
              "thực sự được mua. Thứ tự ghép và kiểu ghép được lựa chọn có chủ đích như Bảng 5.1.")
    table(doc, ["Bước", "Bảng trái", "Bảng phải", "Khóa", "Kiểu ghép", "Lý do"],
          [["1", "order_items", "products", "product_id", "left, N:1",
            "Giữ mọi dòng mua hàng, kể cả sản phẩm thiếu danh mục để xử lý sau"],
           ["2", "kết quả (1)", "translation", "product_category_name", "left, N:1",
            "Không làm mất dòng của danh mục chưa có bản dịch"],
           ["3", "kết quả (2)", "orders", "order_id", "inner, N:1",
            "Chỉ giữ dòng có thông tin đơn hàng, lấy `order_status`"]],
          [1.5, 2.2, 2.1, 3.3, 2.0, 4.4], cap="Bảng 5.1: Chiến lược ghép bảng",
          align=["c", "l", "l", "l", "c", "l"], font_size=10.5)
    para(doc, "Tham số `validate=\"many_to_one\"` của `pandas.merge` buộc bảng phải thỏa quan hệ "
              "nhiều – một: nếu bảng phải có khóa trùng lặp, pandas sẽ báo lỗi ngay thay vì âm "
              "thầm nhân bản dòng. Sau khi ghép, một câu lệnh `assert` đối chiếu số dòng kết quả "
              "với số dòng ban đầu của `order_items` (112.650 dòng). Kết quả: số dòng được bảo "
              "toàn tuyệt đối, chứng tỏ phép ghép đúng khóa.")
    code_block(doc, '''
df = items.merge(products, on="product_id", how="left", validate="many_to_one")
df = df.merge(translation, on="product_category_name", how="left",
              validate="many_to_one")
df = df.merge(orders, on="order_id", how="inner", validate="many_to_one")
assert len(df) == len(items), "Số dòng thay đổi sau merge -> khóa bị nhân bản!"''', cap="Đoạn mã 5.2: Ghép bảng có kiểm chứng quan hệ khóa")

    h2(doc, "5.4. Xử lý giá trị thiếu")
    para(doc, "Nhóm áp dụng ba quy tắc xử lý tương ứng với ba tình huống thiếu dữ liệu ảnh "
              "hưởng tới cột danh mục:")
    table(doc, ["Tình huống", "Cách xử lý", "Lý do lựa chọn"],
          [["Sản phẩm không có danh mục (610 sản phẩm, [[số dòng]] dòng mua hàng)",
            "Gán nhãn `unknown`; giữ lại khi thống kê nhưng **loại khỏi giao dịch**",
            "Không thể suy ra danh mục một cách tin cậy; giữ nhãn giúp minh bạch số lượng bị loại"],
           ["Danh mục chưa có bản dịch (`pc_gamer`, `portateis_cozinha_…`)",
            "Bổ sung bản dịch thủ công: `pc_gamer`, `portable_kitchen_food_processors`",
            "Giữ lại thông tin danh mục thật, tránh mất item"],
           ["Danh mục khác vẫn thiếu bản dịch (nếu có)",
            "Dùng tên tiếng Bồ Đào Nha làm phương án dự phòng",
            "Đảm bảo cột `category` không còn giá trị thiếu"]],
          [4.6, 5.4, 5.5], cap="Bảng 5.2: Quy tắc xử lý giá trị thiếu", font_size=11)
    para(doc, "Nhóm **không** dùng các phương pháp điền khuyết như điền giá trị phổ biến nhất, "
              "vì việc gán sai danh mục sẽ tạo ra các quan hệ mua kèm giả, làm sai lệch trực "
              "tiếp kết quả khai thác. Sau bước này, cột `category` không còn giá trị thiếu. Các "
              "cột thời gian còn thiếu trong `orders` được giữ nguyên vì không tham gia vào bài "
              "toán.")

    h2(doc, "5.5. Lọc đơn hàng theo trạng thái")
    para(doc, "Các đơn có trạng thái `canceled` (đã hủy) và `unavailable` (không khả dụng) bị "
              "loại bỏ vì không phản ánh giao dịch thực sự hoàn tất. Các trạng thái còn lại như "
              "`shipped`, `invoiced`, `processing` vẫn được giữ, vì tại thời điểm đặt hàng khách "
              "đã thể hiện ý định mua đầy đủ các sản phẩm trong giỏ, đây chính là thông tin mà "
              "phân tích giỏ hàng cần. Sau khi lọc, số đơn hàng còn lại là [[số đơn sau lọc]], "
              "tương ứng [[số dòng sau lọc]] dòng sản phẩm.")

    h2(doc, "5.6. Xử lý giá trị ngoại lệ")
    para(doc, "Giá trị ngoại lệ của `price` được phát hiện bằng quy tắc khoảng tứ phân vị (IQR):")
    formula(doc, "IQR = Q₃ − Q₁ ,     ngưỡng trên = Q₃ + 1,5 × IQR")
    para(doc, "Với dữ liệu Olist, ngưỡng trên xấp xỉ [[ngưỡng IQR]] BRL và có khoảng "
              "[[tỷ lệ % dòng ngoại lệ]] số dòng vượt ngưỡng. Tuy nhiên, nhóm quyết định **chỉ "
              "gắn cờ** `is_price_outlier`, **không loại bỏ** các dòng này, vì hai lý do. Thứ "
              "nhất, luật kết hợp chỉ dựa trên sự *có mặt* của danh mục trong giỏ hàng, giá trị "
              "giá không tham gia vào việc tính Support, Confidence hay Lift. Thứ hai, các sản "
              "phẩm giá cao là giao dịch hợp lệ (đồ điện tử, nội thất…); loại bỏ chúng sẽ làm "
              "mất quan hệ mua kèm thật của những danh mục giá trị cao. Cờ ngoại lệ vẫn được lưu "
              "để phục vụ các phân tích mở rộng sau này.")

    h2(doc, "5.7. Tạo giao dịch theo order_id")
    para(doc, "Mỗi đơn hàng được chuyển thành một giao dịch là **tập các danh mục khác nhau** "
              "xuất hiện trong đơn. Có ba quyết định thiết kế quan trọng:")
    bullets(doc, [
        "**Đơn vị giao dịch là order_id**, không phải khách hàng: luật phản ánh hành vi mua "
        "cùng lúc, phù hợp trực tiếp với gợi ý “mua kèm” tại thời điểm thanh toán.",
        "**Khử trùng danh mục trong đơn:** một đơn mua 3 sản phẩm cùng danh mục "
        "`bed_bath_table` chỉ tính danh mục này một lần. Luật kết hợp truyền thống làm việc trên "
        "dữ liệu nhị phân (có/không), số lượng không có ý nghĩa.",
        "**Loại nhãn unknown:** nếu giữ lại, `unknown` sẽ tạo ra các luật vô nghĩa như "
        "“unknown → health_beauty”.",
    ])
    code_block(doc, '''
df = df[df["category"] != UNKNOWN_CATEGORY]
tx = df.groupby("order_id")["category"].apply(lambda s: sorted(set(s)))''', cap="Đoạn mã 5.3: Tạo giao dịch theo order_id")
    para(doc, "Bảng 5.3 minh họa quá trình chuyển đổi từ dữ liệu chi tiết sang giao dịch.")
    table(doc, ["order_id", "Các dòng trong order_items (danh mục)", "Giao dịch"],
          [["A", "bed_bath_table, bed_bath_table, furniture_decor", "{bed_bath_table, furniture_decor}"],
           ["B", "health_beauty", "{health_beauty}"],
           ["C", "baby, toys, baby", "{baby, toys}"],
           ["D", "unknown, sports_leisure", "{sports_leisure}"]],
          [2.0, 7.3, 6.2], cap="Bảng 5.3: Ví dụ minh họa tạo giao dịch (dữ liệu giả định)",
          font_size=11)

    h2(doc, "5.8. Mã hóa one-hot")
    para(doc, "Các thuật toán trong `mlxtend` yêu cầu đầu vào là một DataFrame nhị phân: mỗi "
              "dòng là một giao dịch, mỗi cột là một item, giá trị `True` nếu item có trong giao "
              "dịch. Nhóm dùng lớp `TransactionEncoder` để thực hiện việc này. Với các giao dịch "
              "ở Bảng 5.3, ma trận thu được như Bảng 5.4.")
    table(doc, ["order_id", "baby", "bed_bath_table", "furniture_decor", "health_beauty",
                "sports_leisure", "toys"],
          [["A", "0", "1", "1", "0", "0", "0"],
           ["B", "0", "0", "0", "1", "0", "0"],
           ["C", "1", "0", "0", "0", "0", "1"],
           ["D", "0", "0", "0", "0", "1", "0"]],
          [1.8, 1.6, 2.6, 2.6, 2.4, 2.4, 1.6], cap="Bảng 5.4: Ma trận one-hot tương ứng",
          align=["c"] * 7, font_size=10.5)
    code_block(doc, '''
te = TransactionEncoder()
arr = te.fit(transactions.tolist()).transform(transactions.tolist())
onehot = pd.DataFrame(arr, columns=te.columns_, index=transactions.index)''', cap="Đoạn mã 5.4: Mã hóa one-hot bằng TransactionEncoder")
    para(doc, "Với dữ liệu thật, ma trận one-hot có kích thước [[số giao dịch]] × [[số danh "
              "mục]]. Vì số cột nhỏ, ma trận dạng dày (dense) với kiểu `bool` chỉ chiếm vài MB "
              "bộ nhớ, nên không cần dùng ma trận thưa.")

    h2(doc, "5.9. Phân tích độ thưa và thiết kế hai kịch bản")
    para(doc, "Sau khi tạo giao dịch, nhóm thống kê số danh mục khác nhau trong mỗi giỏ hàng "
              "(Hình 5.2 và Bảng 5.5). Kết quả xác nhận nhận định ở mục 4.6: phần lớn giao dịch "
              "chỉ có **đúng một danh mục**, và mật độ của ma trận one-hot chỉ khoảng "
              "[[mật độ %]].")
    figure(doc, OUT_FIG / "01_basket_distribution.png",
           "Hình 5.2: Phân bố số danh mục trong một đơn hàng và top 15 danh mục",
           hint="outputs/figures/01_basket_distribution.png (sinh bởi phase2_mining.py)")
    table(doc, ["Số danh mục trong đơn", "Số đơn hàng", "Tỷ lệ"],
          [["1", "[[...]]", "[[...]]"], ["2", "[[...]]", "[[...]]"], ["3", "[[...]]", "[[...]]"],
           ["≥ 4", "[[...]]", "[[...]]"]],
          [5.0, 4.0, 3.0], cap="Bảng 5.5: Phân bố kích thước giỏ hàng "
                               "(nguồn: outputs/tables/basket_size_distribution.csv)",
          align=["c", "r", "r"])
    para(doc, "Hệ quả của độ thưa này là rất lớn. Một giao dịch chỉ có một danh mục không bao "
              "giờ đóng góp vào support của bất kỳ tập hai phần tử nào. Do đó, nếu chỉ có khoảng "
              "[[x]]% đơn hàng chứa từ hai danh mục trở lên, thì support của mọi cặp danh mục "
              "đều nhỏ hơn [[x]]%, buộc ngưỡng minsup phải rất thấp (cỡ 10⁻⁴ đến 10⁻³) mới tìm "
              "được luật. Để phân tích đầy đủ, nhóm thiết kế **hai kịch bản** dữ liệu (Bảng 5.6).")
    table(doc, ["Tiêu chí", "Kịch bản ALL", "Kịch bản MULTI"],
          [["Tập giao dịch", "Toàn bộ đơn hàng hợp lệ", "Chỉ đơn có ≥ 2 danh mục khác nhau"],
           ["Số giao dịch", "[[...]]", "[[...]]"],
           ["Dải minsup khảo sát", "10⁻² → 10⁻⁴", "10⁻¹ → 3×10⁻³"],
           ["Ý nghĩa của support", "Tỷ lệ trên **toàn thị trường**",
            "Tỷ lệ trong nhóm khách **đã mua nhiều danh mục**"],
           ["Ưu điểm", "Support, confidence, lift phản ánh đúng thực tế",
            "Mẫu mua kèm nổi bật hơn, minsup dễ chọn hơn"],
           ["Hạn chế", "Support rất nhỏ, cần minsup thấp",
            "Confidence bị thổi phồng do đã loại giỏ một danh mục"],
           ["Vai trò trong đề tài", "**Sinh luật cuối cùng** và gợi ý trong app",
            "So sánh hiệu năng, đối chiếu cấu trúc luật"]],
          [3.6, 6.0, 5.9], cap="Bảng 5.6: Hai kịch bản dữ liệu phục vụ khai thác", font_size=11)
    para(doc, "Luật cuối cùng dùng cho diễn giải và ứng dụng demo được sinh từ kịch bản ALL, "
              "để các độ đo phản ánh đúng xác suất trên toàn bộ khách hàng. Kịch bản MULTI được "
              "dùng làm đối chứng: nếu một luật mạnh ở cả hai kịch bản thì luật đó càng đáng tin "
              "cậy.")

    h2(doc, "5.10. Tổng kết tiền xử lý")
    table(doc, ["Bước", "Trước", "Sau", "Ghi chú"],
          [["Ghép bảng", "112.650 dòng order_items", "112.650 dòng", "Bảo toàn số dòng (assert)"],
           ["Lọc trạng thái", "[[...]] đơn", "[[...]] đơn", "Bỏ canceled, unavailable"],
           ["Xử lý thiếu danh mục", "[[...]] dòng `unknown`", "Loại khỏi giao dịch", "Không điền khuyết"],
           ["Tạo giao dịch", "[[...]] dòng", "[[...]] giao dịch", "Khử trùng danh mục trong đơn"],
           ["One-hot", "[[...]] giao dịch", "Ma trận [[...]] × [[...]]", "Kiểu bool, dạng dày"]],
          [3.4, 4.0, 3.8, 4.3], cap="Bảng 5.7: Tóm tắt kết quả các bước tiền xử lý", font_size=11)
    para(doc, "Sau tiền xử lý, dữ liệu đã sẵn sàng cho bước khai thác: mỗi giao dịch là một tập "
              "danh mục sạch, đã chuẩn hóa tên tiếng Anh, không chứa danh mục không xác định và "
              "không chứa đơn hàng bị hủy. Các bảng trung gian đều được lưu ra tệp để đảm bảo "
              "khả năng tái lập kết quả.")


# ============================================================================
def build():
    REPORT_FIG.mkdir(parents=True, exist_ok=True)
    draw_schema(REPORT_FIG / "schema.png")
    draw_pipeline(REPORT_FIG / "pipeline.png")

    doc = new_document()
    cover(doc)

    sec = new_section(doc)          # phần mở đầu: đánh số La Mã i, ii, iii
    add_page_number_footer(sec, start_at=1, fmt="lowerRoman")
    assignment(doc)
    page_break(doc)
    add_toc(doc)
    page_break(doc)
    abbreviations(doc)

    sec = new_section(doc)          # nội dung chính: đánh số lại từ 1
    add_page_number_footer(sec, start_at=1)
    section3(doc)
    page_break(doc)
    section4(doc)
    page_break(doc)
    section5(doc)

    doc.save(OUT_FILE)
    print(f"Đã tạo {OUT_FILE}")


if __name__ == "__main__":
    build()
